"""Load supported source files into LangChain documents."""

from pathlib import Path
import logging
import time

from bs4 import BeautifulSoup
from langchain_community.document_loaders import Docx2txtLoader, PyMuPDFLoader, WebBaseLoader
from langchain_core.documents import Document
import openpyxl


logger = logging.getLogger(__name__)


def load_excel(file_path: str) -> list[Document]:
    """Convert spreadsheet rows into short natural-language documents."""

    # data_only=True 读取公式计算结果；read_only=True 避免大表构建完整工作簿对象。
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        return []

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    # 优先按中文表头识别字段，识别失败时兼容第一阶段约定的第 2、3 列。
    name_column = next(
        (index for index, header in enumerate(headers) if any(key in header for key in ("名称", "商品", "品名"))),
        1,
    )
    price_column = next(
        (index for index, header in enumerate(headers) if any(key in header for key in ("单价", "价格", "价"))),
        2,
    )

    documents: list[Document] = []
    item_names: list[str] = []
    filename = Path(file_path).name

    for row_index, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        name = row[name_column] if name_column < len(row) else None
        price = row[price_column] if price_column < len(row) else None
        if name is None or not str(name).strip():
            continue

        item_name = str(name).strip()
        if price is None or not str(price).strip():
            text = f"{item_name}：暂无价格信息。"
        else:
            try:
                numeric_price = float(price)
                price_text = (
                    str(int(numeric_price))
                    if numeric_price == int(numeric_price)
                    else str(round(numeric_price, 2))
                )
            except (TypeError, ValueError):
                price_text = str(price).strip()
            text = f"{item_name}的单价是{price_text}元。"

        item_names.append(item_name)
        # 每一行作为独立语义单元，保留原始行号，回答价格问题时可精确引用。
        documents.append(
            Document(
                page_content=text,
                metadata={"source": filename, "row": row_index, "type": "price_table"},
            )
        )

    if item_names:
        summary = f"价格表【{filename}】共包含以下商品：\n" + "、".join(item_names[:100])
        if len(item_names) > 100:
            summary += f"……等共 {len(item_names)} 种商品。"
        documents.insert(
            0,
            Document(
                page_content=summary,
                metadata={"source": filename, "type": "price_summary"},
            ),
        )
    return documents


def load_html(file_path: str) -> list[Document]:
    """Extract visible text from a local HTML file.

    The first version intentionally ignores JavaScript-rendered content and
    linked pages. It only indexes text present in the uploaded HTML file.
    """

    path = Path(file_path)
    # 直接传入字节，让 BeautifulSoup 根据 HTML 声明推断编码，兼容常见中文页面。
    soup = BeautifulSoup(path.read_bytes(), "html.parser")
    # 脚本、样式和模板不是用户可见正文，写入向量库会制造检索噪声。
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()

    title = soup.title.get_text(" ", strip=True) if soup.title else None
    lines = [line.strip() for line in soup.get_text("\n").splitlines()]
    text = "\n".join(line for line in lines if line)
    if not text:
        raise ValueError("HTML 文件中没有可索引的文本内容")

    metadata = {"source": path.name, "type": "html"}
    if title:
        metadata["title"] = title
    return [Document(page_content=text, metadata=metadata)]


def load_file(file_path: str) -> list[Document]:
    started = time.perf_counter()
    # FastAPI 当前只开放本地上传；URL 分支仅用于兼容第一阶段已有调用。
    if file_path.startswith(("http://", "https://")):
        documents = WebBaseLoader(file_path).load()
        source_type = "web"
    else:
        extension = Path(file_path).suffix.lower()
        if extension == ".pdf":
            documents = PyMuPDFLoader(file_path).load()
            source_type = "pdf"
        elif extension == ".docx":
            documents = Docx2txtLoader(file_path).load()
            source_type = "docx"
        elif extension in {".xlsx", ".xls"}:
            documents = load_excel(file_path)
            source_type = "spreadsheet"
        elif extension in {".html", ".htm"}:
            documents = load_html(file_path)
            source_type = "html"
        else:
            raise ValueError(f"不支持的格式: {extension or '未知'}")

    logger.info(
        "document loaded type=%s documents=%d latency_ms=%d",
        source_type,
        len(documents),
        int((time.perf_counter() - started) * 1000),
    )
    return documents
